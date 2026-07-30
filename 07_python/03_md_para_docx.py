# -*- coding: utf-8 -*-
"""Converte os documentos Markdown do projeto para DOCX formatado (SEDEC/EUROCLIMA+)."""
import re, sys, os
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

AZUL = RGBColor(0x1F, 0x5C, 0x8B)
CINZA = RGBColor(0x44, 0x44, 0x44)


def sombrear(cell, hexcor):
    tc = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), hexcor)
    tc.append(sh)


def inline(par, texto):
    """Aplica **negrito**, *itálico* e `código`."""
    for parte in re.split(r"(\*\*.+?\*\*|(?<!\*)\*[^*]+?\*(?!\*)|`[^`]+?`)", texto):
        if not parte:
            continue
        if parte.startswith("**") and parte.endswith("**"):
            par.add_run(parte[2:-2]).bold = True
        elif parte.startswith("`") and parte.endswith("`"):
            r = par.add_run(parte[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(9)
        elif parte.startswith("*") and parte.endswith("*"):
            par.add_run(parte[1:-1]).italic = True
        else:
            par.add_run(parte)


def converte(md_path, docx_path, titulo=None):
    linhas = open(md_path, encoding="utf-8").read().split("\n")
    doc = docx.Document()

    st = doc.styles["Normal"]
    st.font.name = "Calibri"; st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.5)
        s.top_margin = s.bottom_margin = Cm(2.2)

    i = 0
    while i < len(linhas):
        ln = linhas[i]
        s = ln.strip()

        # --------------------------------------------------------- tabela
        if s.startswith("|") and i + 1 < len(linhas) and re.match(r"^\|[\s:\-|]+\|$", linhas[i + 1].strip()):
            cab = [c.strip() for c in s.strip("|").split("|")]
            i += 2
            corpo = []
            while i < len(linhas) and linhas[i].strip().startswith("|"):
                corpo.append([c.strip() for c in linhas[i].strip().strip("|").split("|")])
                i += 1
            t = doc.add_table(rows=1, cols=len(cab))
            t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for k, c in enumerate(cab):
                cel = t.rows[0].cells[k]
                cel.text = ""
                p = cel.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                inline(p, c)
                for r in p.runs:
                    r.bold = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                sombrear(cel, "1F5C8B")
            for row in corpo:
                cells = t.add_row().cells
                for k in range(min(len(row), len(cab))):
                    cells[k].text = ""
                    p = cells[k].paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    inline(p, row[k].replace("<br>", " "))
                    for r in p.runs:
                        r.font.size = Pt(9)
            doc.add_paragraph()
            continue

        # ------------------------------------------------------- títulos
        if s.startswith("#"):
            n = len(s) - len(s.lstrip("#"))
            txt = s[n:].strip()
            h = doc.add_heading(level=min(n, 4))
            h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            inline(h, txt)
            for r in h.runs:
                r.font.color.rgb = AZUL
                r.font.name = "Calibri"
            i += 1
            continue

        # ------------------------------------------------------ separador
        if s in ("---", "***", "___"):
            doc.add_paragraph()
            i += 1
            continue

        # ------------------------------------------------------ citação
        if s.startswith(">"):
            bloco = []
            while i < len(linhas) and linhas[i].strip().startswith(">"):
                bloco.append(linhas[i].strip().lstrip(">").strip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.space_before = Pt(6)
            inline(p, " ".join(x for x in bloco if x))
            for r in p.runs:
                r.font.size = Pt(9.5); r.font.color.rgb = CINZA
            continue

        # -------------------------------------------------------- listas
        m = re.match(r"^[-*]\s+(.*)", s)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            inline(p, m.group(1))
            i += 1
            continue
        m = re.match(r"^(\d+)\.\s+(.*)", s)
        if m:
            p = doc.add_paragraph(style="List Number")
            inline(p, m.group(2))
            i += 1
            continue

        # ------------------------------------------------------ parágrafo
        if s:
            p = doc.add_paragraph()
            inline(p, s)
        i += 1

    doc.save(docx_path)
    print(f"gerado: {docx_path}")


if __name__ == "__main__":
    base = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    alvos = [
        (os.path.join(base, "06_resultados", "TR_MINUTA_REV0B.md"),
         os.path.join(base, "06_resultados", "TR_MINUTA_REV0B.docx")),
        (os.path.join(base, "06_resultados", "NOTA_TECNICA_PRELIMINAR_BARRAGENS.md"),
         os.path.join(base, "06_resultados", "NOTA_TECNICA_PRELIMINAR_BARRAGENS.docx")),
        (os.path.join(base, "PLANO_DE_ACAO.md"),
         os.path.join(base, "PLANO_DE_ACAO.docx")),
    ]
    for md, dx in alvos:
        if os.path.exists(md):
            converte(md, dx)
